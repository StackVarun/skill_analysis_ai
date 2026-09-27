"""Tests for resume upload APIs."""
import pytest
import json
import io
from docx import Document


class TestResume:
    """Test resume upload endpoints."""

    def create_test_docx(self):
        """Create a test DOCX file in memory."""
        doc = Document()
        doc.add_heading('John Doe', 0)
        doc.add_heading('Software Engineer', level=1)
        doc.add_paragraph('Email: john.doe@example.com')
        doc.add_paragraph('Phone: +1-555-123-4567')
        doc.add_heading('Experience', level=1)
        doc.add_paragraph('Tech Corp - Software Engineer (2022-2024)')
        doc.add_heading('Skills', level=1)
        doc.add_paragraph('Python, JavaScript, React, SQL, Flask, PostgreSQL')
        
        # Save to bytes
        file_stream = io.BytesIO()
        doc.save(file_stream)
        file_stream.seek(0)
        return file_stream

    def test_list_resumes_empty(self, client, auth_headers, student_profile_id):
        """Test listing resumes when none exist."""
        response = client.get("/api/students/resume", headers=auth_headers)
        assert response.status_code == 200
        data = json.loads(response.data)
        assert data == []

    def test_upload_resume_docx(self, client, auth_headers, student_profile_id):
        """Test uploading a DOCX resume."""
        file_stream = self.create_test_docx()
        
        data = {"file": (file_stream, "test_resume.docx")}
        response = client.post("/api/students/resume",
                             data=data,
                             headers=auth_headers,
                             content_type="multipart/form-data")
        assert response.status_code == 201
        data = json.loads(response.data)
        assert data["file_type"] == "docx"
        assert data["status"] == "COMPLETED"
        assert data["extracted_text_length"] > 0

    def test_upload_resume_invalid_file_type(self, client, auth_headers, student_profile_id):
        """Test uploading an invalid file type."""
        file_stream = io.BytesIO(b"test content")
        data = {"file": (file_stream, "test_resume.txt")}
        response = client.post("/api/students/resume",
                             data=data,
                             headers=auth_headers,
                             content_type="multipart/form-data")
        assert response.status_code == 400
        data = json.loads(response.data)
        assert "Invalid file type" in data["message"]

    def test_upload_resume_no_file(self, client, auth_headers, student_profile_id):
        """Test uploading without a file."""
        response = client.post("/api/students/resume",
                             data={},
                             headers=auth_headers,
                             content_type="multipart/form-data")
        assert response.status_code == 400
        data = json.loads(response.data)
        assert "Please provide a resume file" in data["message"]

    def test_get_resume(self, client, auth_headers, student_profile_id):
        """Test getting a specific resume."""
        file_stream = self.create_test_docx()
        data = {"file": (file_stream, "test_resume.docx")}
        create_resp = client.post("/api/students/resume",
                                data=data,
                                headers=auth_headers,
                                content_type="multipart/form-data")
        resume_id = json.loads(create_resp.data)["id"]
        
        response = client.get(f"/api/students/resume/{resume_id}", headers=auth_headers)
        assert response.status_code == 200
        data = json.loads(response.data)
        assert data["id"] == resume_id

    def test_delete_resume(self, client, auth_headers, student_profile_id):
        """Test deleting a resume."""
        file_stream = self.create_test_docx()
        data = {"file": (file_stream, "test_resume.docx")}
        create_resp = client.post("/api/students/resume",
                                data=data,
                                headers=auth_headers,
                                content_type="multipart/form-data")
        resume_id = json.loads(create_resp.data)["id"]
        
        response = client.delete(f"/api/students/resume/{resume_id}", headers=auth_headers)
        assert response.status_code == 200
        data = json.loads(response.data)
        assert data["message"] == "Resume deleted successfully"
        
        # Verify it's gone
        response = client.get("/api/students/resume", headers=auth_headers)
        data = json.loads(response.data)
        assert len(data) == 0

    def test_resume_authorization_industry(self, client, industry_auth_headers):
        """Test that industry users cannot access resumes."""
        response = client.get("/api/students/resume", headers=industry_auth_headers)
        assert response.status_code == 403