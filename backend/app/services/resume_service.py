"""Resume service for business logic."""
import os
import uuid
import fitz  # PyMuPDF
from docx import Document
from app.extensions import db
from app.models.student_profile import StudentProfile
from app.models.resume import Resume, ResumeStatus, UPLOAD_FOLDER, ALLOWED_EXTENSIONS, MAX_FILE_SIZE


class ResumeService:
    """Service layer for resume operations."""

    @staticmethod
    def get_student_resumes(student_id: int):
        """Get all resumes for a student."""
        return Resume.query.filter_by(student_id=student_id).order_by(Resume.created_at.desc()).all()

    @staticmethod
    def get_resume_by_id(resume_id: int):
        """Get resume by ID."""
        return Resume.query.filter_by(id=resume_id).first()

    @staticmethod
    def get_student_resume(student_id: int, resume_id: int):
        """Get a specific student's resume."""
        return Resume.query.filter_by(student_id=student_id, id=resume_id).first()

    @staticmethod
    def allowed_file(filename: str) -> bool:
        """Check if file extension is allowed."""
        return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS

    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        """Extract text from PDF file using PyMuPDF."""
        text = ""
        try:
            doc = fitz.open(file_path)
            for page_num in range(len(doc)):
                page = doc[page_num]
                text += page.get_text()
            doc.close()
        except Exception as e:
            raise ValueError(f"Failed to extract text from PDF: {str(e)}")
        return text

    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        """Extract text from DOCX file using python-docx."""
        text = ""
        try:
            doc = Document(file_path)
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + "\n"
        except Exception as e:
            raise ValueError(f"Failed to extract text from DOCX: {str(e)}")
        return text

    @staticmethod
    def extract_text(file_path: str, file_type: str) -> str:
        """Extract text based on file type."""
        if file_type == "pdf":
            return ResumeService.extract_text_from_pdf(file_path)
        elif file_type == "docx":
            return ResumeService.extract_text_from_docx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")

    @staticmethod
    def save_uploaded_file(file, student_id: int) -> tuple:
        """Save uploaded file and return (file_path, filename, file_type, file_size)."""
        filename = file.filename
        if not filename or not ResumeService.allowed_file(filename):
            raise ValueError("Invalid file type. Only PDF and DOCX files are allowed.")

        file_type = filename.rsplit(".", 1)[1].lower()

        # Check file size
        file.seek(0, os.SEEK_END)
        file_size = file.tell()
        file.seek(0)

        if file_size > MAX_FILE_SIZE:
            raise ValueError(f"File size exceeds maximum allowed size of {MAX_FILE_SIZE} bytes")

        # Create upload directory if it doesn't exist
        os.makedirs(UPLOAD_FOLDER, exist_ok=True)

        # Generate unique filename
        unique_filename = f"{student_id}_{uuid.uuid4().hex}_{filename}"
        file_path = os.path.join(UPLOAD_FOLDER, unique_filename)

        # Sanitize file path
        file_path = os.path.normpath(file_path)
        if not file_path.startswith(os.path.abspath(UPLOAD_FOLDER)):
            raise ValueError("Invalid file path")

        # Save file
        file.save(file_path)

        return file_path, unique_filename, file_type, file_size

    @staticmethod
    def create_resume(student_id: int, file) -> Resume:
        """Create a new resume record and process the uploaded file."""
        student = StudentProfile.query.filter_by(id=student_id).first()
        if not student:
            raise ValueError("Student profile not found")

        # Save file
        file_path, filename, file_type, file_size = ResumeService.save_uploaded_file(file, student_id)

        # Create resume record
        resume = Resume(
            student_id=student_id,
            filename=filename,
            file_type=file_type,
            file_path=file_path,
            file_size=file_size,
            status=ResumeStatus.PROCESSING
        )

        db.session.add(resume)
        db.session.commit()

        # Extract text in background (for now, do it synchronously)
        try:
            extracted_text = ResumeService.extract_text(file_path, file_type)
            resume.extracted_text = extracted_text
            resume.extracted_text_length = len(extracted_text)
            resume.status = ResumeStatus.COMPLETED
        except Exception as e:
            resume.status = ResumeStatus.FAILED
            resume.error_message = str(e)

        db.session.commit()
        return resume

    @staticmethod
    def delete_resume(resume_id: int, student_id: int) -> bool:
        """Delete a student's resume and the associated file."""
        resume = Resume.query.filter_by(id=resume_id, student_id=student_id).first()
        if not resume:
            return False

        # Delete file if it exists
        if os.path.exists(resume.file_path):
            try:
                os.remove(resume.file_path)
            except OSError:
                pass  # File might already be deleted

        db.session.delete(resume)
        db.session.commit()
        return True